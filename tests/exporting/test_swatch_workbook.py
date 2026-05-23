from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook
from PIL import Image

from src.common.errors import ExportError
from src.exporting import export_swatch_workbook
from src.extraction.extractor import ExtractedFieldValue, ExtractedGroup


def _png_bytes(color: tuple[int, int, int] = (20, 120, 80)) -> bytes:
    image = Image.new("RGB", (32, 24), color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_export_swatch_workbook_writes_dynamic_headers_and_rows(tmp_path) -> None:
    path = tmp_path / "swatches.xlsx"
    groups = [
        ExtractedGroup(
            values={
                "swatch": ExtractedFieldValue("image", image_bytes=_png_bytes()),
                "material_id": ExtractedFieldValue("text", text="MAT-001"),
                "description": ExtractedFieldValue("text", text="Green"),
            }
        ),
        ExtractedGroup(
            values={
                "swatch": ExtractedFieldValue("image", image_bytes=_png_bytes((40, 80, 180))),
                "material_id": ExtractedFieldValue("text", text="MAT-002"),
                "description": ExtractedFieldValue("text", text="Blue"),
            }
        ),
    ]

    export_swatch_workbook(
        path,
        groups,
        manufacturer="Acme Finishes",
        category="Fabric",
    )

    workbook = load_workbook(path)
    sheet = workbook["Extracted Data"]

    assert sheet["A1"].value == "Manufacturer"
    assert sheet["B1"].value == "Acme Finishes"
    assert sheet["A2"].value == "Category"
    assert sheet["B2"].value == "Fabric"
    assert sheet["A4"].value == "swatch"
    assert sheet["B4"].value == "material_id"
    assert sheet["C4"].value == "description"
    assert sheet["B5"].value == "MAT-001"
    assert sheet["C6"].value == "Blue"
    assert len(sheet._images) == 2


def test_export_swatch_workbook_skips_invalid_image_bytes(tmp_path) -> None:
    path = tmp_path / "swatches.xlsx"
    groups = [
        ExtractedGroup(
            values={
                "swatch": ExtractedFieldValue("image", image_bytes=b"not an image"),
                "material_id": ExtractedFieldValue("text", text="MAT-001"),
            }
        )
    ]

    export_swatch_workbook(path, groups, manufacturer="", category="")

    workbook = load_workbook(path)
    sheet = workbook["Extracted Data"]

    assert sheet["B5"].value == "MAT-001"
    assert sheet._images == []


def test_export_swatch_workbook_wraps_save_failure(monkeypatch, tmp_path) -> None:
    path = tmp_path / "swatches.xlsx"
    groups = [
        ExtractedGroup(
            values={
                "material_id": ExtractedFieldValue("text", text="MAT-001"),
            }
        )
    ]

    def fail_save(_self, _path) -> None:
        raise PermissionError("file is locked")

    monkeypatch.setattr("src.exporting.swatch_workbook.Workbook.save", fail_save)

    with pytest.raises(ExportError) as exc_info:
        export_swatch_workbook(path, groups, manufacturer="", category="")

    assert "Could not export to Excel" in exc_info.value.user_message
