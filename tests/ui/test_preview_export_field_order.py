from __future__ import annotations

import io
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from PyQt6.QtWidgets import QApplication

from src.exporting import export_swatch_workbook
from src.extraction.extractor import ExtractedFieldValue, ExtractedGroup
from src.ui.panels.preview_panel import PreviewPanel


def _png_bytes(color: tuple[int, int, int]) -> bytes:
    image = Image.new("RGB", (24, 18), color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def test_preview_and_export_use_same_first_seen_field_order(
    qapp: QApplication,
    tmp_path: Path,
) -> None:
    assert qapp is not None
    groups = [
        ExtractedGroup(
            values={
                "material_id": ExtractedFieldValue("text", text="MAT-001"),
                "swatch": ExtractedFieldValue("image", image_bytes=_png_bytes((20, 140, 80))),
            }
        ),
        ExtractedGroup(
            values={
                "description": ExtractedFieldValue("text", text="Satin"),
                "material_id": ExtractedFieldValue("text", text="MAT-002"),
                "swatch": ExtractedFieldValue("image", image_bytes=_png_bytes((50, 90, 170))),
            }
        ),
    ]

    panel = PreviewPanel()
    panel.load(groups)

    preview_headers: list[str] = []
    for i in range(panel._table.columnCount()):  # noqa: SLF001 - contract test for header ordering
        item = panel._table.horizontalHeaderItem(i)  # noqa: SLF001 - contract test for header ordering
        assert item is not None
        preview_headers.append(item.text())

    path = tmp_path / "field-order.xlsx"
    export_swatch_workbook(path, groups, manufacturer="Acme", category="Fabric")
    sheet = load_workbook(path)["Extracted Data"]
    export_headers = [
        sheet.cell(row=4, column=1).value,
        sheet.cell(row=4, column=2).value,
        sheet.cell(row=4, column=3).value,
    ]

    assert preview_headers == ["material_id", "swatch", "description"]
    assert export_headers == preview_headers


