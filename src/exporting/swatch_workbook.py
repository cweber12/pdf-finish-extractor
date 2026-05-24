from __future__ import annotations

import io
from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

from src.common.errors import ExportError
from src.extraction.extractor import ExtractedGroup
from src.extraction.group_projection import projected_field_names

_HEADER_FILL = PatternFill("solid", fgColor="E2E8F0")
_COLUMN_FILL = PatternFill("solid", fgColor="CBD5E1")
_BOLD = Font(name="Arial", bold=True)
_BODY = Font(name="Arial")
_MAX_IMAGE_WIDTH = 140
_MAX_IMAGE_HEIGHT = 92


def export_swatch_workbook(
    path: str | Path,
    groups: Sequence[ExtractedGroup],
    *,
    manufacturer: str,
    category: str,
) -> None:
    """Write extracted groups to an Excel workbook."""
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Extracted Data"

        sheet["A1"] = "Manufacturer"
        sheet["B1"] = manufacturer
        sheet["A2"] = "Category"
        sheet["B2"] = category

        field_names = projected_field_names(groups)
        for column_index, field_name in enumerate(field_names, start=1):
            cell = sheet.cell(row=4, column=column_index, value=field_name)
            cell.font = _BOLD
            cell.fill = _COLUMN_FILL
            cell.alignment = Alignment(horizontal="center")
            sheet.column_dimensions[get_column_letter(column_index)].width = 24

        for cell in ("A1", "A2"):
            sheet[cell].font = _BOLD
        for cell in ("A1", "B1", "A2", "B2"):
            sheet[cell].fill = _HEADER_FILL

        sheet.freeze_panes = "A5"

        for row_index, group in enumerate(groups, start=5):
            sheet.row_dimensions[row_index].height = 74
            for column_index, field_name in enumerate(field_names, start=1):
                value = group.values.get(field_name)
                if value is None:
                    continue
                cell = sheet.cell(row=row_index, column=column_index)
                cell.font = _BODY
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if value.field_type == "image":
                    image = _worksheet_image(value.image_bytes)
                    if image is not None:
                        sheet.add_image(image, f"{get_column_letter(column_index)}{row_index}")
                else:
                    cell.value = value.text

        workbook.save(path)
    except ExportError:
        raise
    except Exception as exc:  # noqa: BLE001 - wrapped to typed export error
        raise ExportError(
            "Could not export to Excel. Close the file if it is open and try again.",
            detail=str(exc),
        ) from exc

def _worksheet_image(image_bytes: bytes) -> WorksheetImage | None:
    try:
        with PILImage.open(io.BytesIO(image_bytes)) as source:
            width, height = _fit_size(source.width, source.height)
            buffer = io.BytesIO()
            source.save(buffer, format="PNG")
    except Exception:
        return None

    buffer.seek(0)
    image = WorksheetImage(buffer)
    image.width = width
    image.height = height
    return image


def _fit_size(width: int, height: int) -> tuple[int, int]:
    if width <= 0 or height <= 0:
        return _MAX_IMAGE_WIDTH, _MAX_IMAGE_HEIGHT

    scale = min(_MAX_IMAGE_WIDTH / width, _MAX_IMAGE_HEIGHT / height, 1)
    return max(1, int(width * scale)), max(1, int(height * scale))
